"""Render a report's Markdown into downloadable formats.

- **md**   — the raw report Markdown.
- **html** — standalone, styled HTML document.
- **pdf**  — HTML rendered to PDF via xhtml2pdf (pure-Python, Windows-friendly).
- **docx** — built with python-docx from a lightweight Markdown parser that handles
  the constructs our reports use: headings, paragraphs, bullet/numbered lists, and
  pipe tables.
"""
from __future__ import annotations

import io
import re

import markdown as md_lib

EXPORT_FORMATS = {
    "md": ("text/markdown", "md"),
    "html": ("text/html", "html"),
    "pdf": ("application/pdf", "pdf"),
    "docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "docx",
    ),
}


class ExportError(RuntimeError):
    pass


_HTML_CSS = """
@page { size: A4; margin: 2cm; }
body { font-family: Helvetica, Arial, sans-serif; color: #1e293b; line-height: 1.5;
       font-size: 11pt; }
h1 { font-size: 20pt; color: #0f172a; }
h2 { font-size: 15pt; color: #0f172a; border-bottom: 1px solid #cbd5e1;
     padding-bottom: 3px; margin-top: 18px; }
h3 { font-size: 12pt; }
table { border-collapse: collapse; width: 100%; margin: 10px 0; font-size: 9pt; }
th, td { border: 1px solid #cbd5e1; padding: 5px 7px; text-align: left; }
th { background: #f1f5f9; }
a { color: #4f46e5; }
code { background: #f1f5f9; padding: 1px 4px; }
"""


def export_report(markdown_text: str, *, title: str, fmt: str) -> bytes:
    if fmt not in EXPORT_FORMATS:
        raise ExportError(f"Unsupported format: {fmt}")
    if not markdown_text:
        raise ExportError("Report has no content to export")

    if fmt == "md":
        return markdown_text.encode("utf-8")
    if fmt == "html":
        return _to_html(markdown_text, title).encode("utf-8")
    if fmt == "pdf":
        return _to_pdf(markdown_text, title)
    if fmt == "docx":
        return _to_docx(markdown_text, title)
    raise ExportError(f"Unsupported format: {fmt}")  # pragma: no cover


def _to_html(markdown_text: str, title: str) -> str:
    body = md_lib.markdown(markdown_text, extensions=["tables", "fenced_code"])
    return (
        f"<!doctype html><html><head><meta charset='utf-8'>"
        f"<title>{_escape(title)}</title><style>{_HTML_CSS}</style></head>"
        f"<body>{body}</body></html>"
    )


def _to_pdf(markdown_text: str, title: str) -> bytes:
    from xhtml2pdf import pisa

    html = _to_html(markdown_text, title)
    buf = io.BytesIO()
    result = pisa.CreatePDF(html, dest=buf, encoding="utf-8")
    if result.err:
        raise ExportError("PDF generation failed")
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# DOCX — minimal Markdown -> python-docx conversion
# --------------------------------------------------------------------------- #
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_ULIST_RE = re.compile(r"^\s*[-*]\s+(.*)$")
_OLIST_RE = re.compile(r"^\s*\d+\.\s+(.*)$")
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}.*$")
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _to_docx(markdown_text: str, title: str) -> bytes:
    from docx import Document

    doc = Document()
    lines = markdown_text.splitlines()
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Table: a header row followed by a separator row of dashes.
        if (
            stripped.startswith("|")
            and i + 1 < n
            and _TABLE_SEP_RE.match(lines[i + 1])
        ):
            i = _emit_table(doc, lines, i)
            continue

        h = _HEADING_RE.match(stripped)
        if h:
            level = min(len(h.group(1)), 4)
            doc.add_heading(_strip_md(h.group(2)), level=level)
            i += 1
            continue

        m = _ULIST_RE.match(line)
        if m:
            _add_rich_paragraph(doc, m.group(1), style="List Bullet")
            i += 1
            continue

        m = _OLIST_RE.match(line)
        if m:
            _add_rich_paragraph(doc, m.group(1), style="List Number")
            i += 1
            continue

        _add_rich_paragraph(doc, stripped)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _emit_table(doc, lines: list[str], start: int) -> int:
    def cells(row: str) -> list[str]:
        return [c.strip() for c in row.strip().strip("|").split("|")]

    header = cells(lines[start])
    i = start + 2  # skip header + separator
    rows: list[list[str]] = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append(cells(lines[i]))
        i += 1

    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for j, text in enumerate(header):
        table.rows[0].cells[j].text = _strip_md(text)
    for r in rows:
        row_cells = table.add_row().cells
        for j in range(len(header)):
            row_cells[j].text = _strip_md(r[j]) if j < len(r) else ""
    doc.add_paragraph("")
    return i


def _add_rich_paragraph(doc, text: str, *, style: str | None = None) -> None:
    """Add a paragraph, rendering **bold** spans as bold runs."""
    para = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    pos = 0
    for match in _BOLD_RE.finditer(text):
        if match.start() > pos:
            para.add_run(text[pos : match.start()])
        para.add_run(match.group(1)).bold = True
        pos = match.end()
    if pos < len(text):
        para.add_run(text[pos:])


def _strip_md(text: str) -> str:
    return _BOLD_RE.sub(r"\1", text).replace("\\|", "|").strip()


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
